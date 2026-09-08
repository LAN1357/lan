from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.reports.management import export_management
from src.workbench.attribution import assign
from src.workbench.closing import close_month, reopen_month
from src.workbench.db import records
from tests.test_workbench_closing import seed
from tests.test_workbench_imports import db


def test_formal_export_requires_close_and_historical_snapshot(db):
    seed(db)
    with pytest.raises(ValueError):
        export_management(db, '2026-09')
    close_month(db, '2026-09', confirmed=True)
    filename, raw = export_management(db, '2026-09')
    assert filename.endswith('_V1.xlsx')
    wb = load_workbook(BytesIO(raw), data_only=False)
    assert len(wb.sheetnames) == 7
    pnl = {r[0].value: r for r in wb['场次损益'].iter_rows()}
    assert pnl['经营复盘利润'][1].value == 1820
    assert pnl['最终结算利润'][1].value == 1120
    assert pnl['最终结算利润'][3].value == sum(c.value for c in pnl['最终结算利润'][1:3])
    assert pnl['结算销售额/投流费'][2].value == '不适用'
    assert pnl['最终结算利润率'][1].number_format == '0.0%'
    assert any('2026-10' in str(r[-1].value) for r in wb['核对明细'])
    reopen_month(db, '2026-09', reason='费用权重更新')
    assign(db, records(db, 'monthly')[0]['id'], {'S1': 1, 'S2': 1}, basis='共享受益', reason='重分配')
    with pytest.raises(ValueError):
        export_management(db, '2026-09')
    _, old_raw = export_management(db, '2026-09', version=1)
    old = load_workbook(BytesIO(old_raw))
    assert list(old['场次损益'].values) == list(wb['场次损益'].values)
    close_month(db, '2026-09', confirmed=True)
    assert export_management(db, '2026-09')[0].endswith('_V2.xlsx')


def test_source_text_is_not_excel_formula(db):
    seed(db)
    assign(db, records(db, 'ads')[0]['id'], {'S1': 1}, basis='=HYPERLINK("https://example.org")', reason='文本安全测试')
    close_month(db, '2026-09', confirmed=True)
    _, raw = export_management(db, '2026-09')
    wb = load_workbook(BytesIO(raw), data_only=False)
    assert any(c.value == '=HYPERLINK("https://example.org")' for row in wb['分摊明细'] for c in row)
    assert all(c.data_type != 'f' for sheet in wb for row in sheet for c in row)
