import json
import sqlite3
from io import BytesIO
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from src.workbench.db import connect, records, dumps
from src.workbench.attribution import assignments, assign, assign_orders
from src.workbench.batches import prepare_assignments, apply_batch, prepare_cost_edit, prepare_cost_fill, prepare_allocation_reuse
from src.workbench.closing import close_month, get_close, preview_close, reopen_month
from src.workbench.costs import save_cost, save_standard
from src.workbench.recommendations import generate, save_ad_link
from src.workbench.workflow import order_queue, fee_states
from src.workbench.imports import import_workbook, SCHEMAS
from src.reports.management import export_management
from tests.test_workbench_imports import db, workbook_bytes
from tests.test_workbench_closing import seed
from tests.test_workbench_m2 import unassign


def test_overlapping_sessions_and_unconfirmed_scope_never_auto_adopt(db):
    ids = unassign(db)
    assert generate(db,ids,session_ids=['S1','S2'])[0]['data']['issues']
    raw = workbook_bytes({'场次':[['S3','丙','2026-09-01 10:00','2026-09-01 12:00']]})
    import_workbook(db,raw,'并行直播.xlsx')
    rec = generate(db,[ids[0]],session_ids=['S1','S3'],source_scope_confirmed=True)[0]
    assert rec['data']['candidates'] == ['S1','S3']
    assert any('多个场次' in x for x in rec['data']['issues'])
    assert ids[0] not in assignments(db)


def test_cross_midnight_time_rule_uses_full_timestamp(db):
    ids = unassign(db)
    session = records(db,'sessions')[0]
    order = records(db,'orders')[0]
    db.execute('UPDATE records SET data=? WHERE id=?',(dumps(session['data']|{'start':'2026-08-31 23:00:00','end':'2026-09-01 01:00:00'}),session['id']))
    db.execute('UPDATE records SET data=? WHERE id=?',(dumps(order['data']|{'paid_at':'2026-09-01 00:15:00'}),order['id']))
    rec = generate(db,[ids[0]],session_ids=['S1'],source_scope_confirmed=True)[0]
    assert rec['data']['targets']=={'S1':1}


def test_write_failure_rolls_back_assignments_and_audits(db):
    ids = unassign(db)
    p = prepare_assignments(db,ids,session_id='S1')
    db.execute(f"CREATE TRIGGER fail_second BEFORE INSERT ON assignments WHEN NEW.record_id={ids[1]} BEGIN SELECT RAISE(ABORT,'simulated disk/write failure'); END")
    with pytest.raises(sqlite3.IntegrityError):
        apply_batch(db,p,reason='批量确认')
    assert all(i not in assignments(db) for i in ids)
    assert db.execute('SELECT COUNT(*) FROM changes WHERE operation_id=?',(p['operation_id'],)).fetchone()[0] == 0


def test_month_lock_after_preview_rejects_all_writes(db):
    seed(db)
    ids = [r['id'] for r in records(db,'orders')]
    p = prepare_assignments(db,ids,session_id='S1',allow_overwrite=True)
    before = assignments(db)
    close_month(db,'2026-09',confirmed=True)
    with pytest.raises(ValueError,match='重开'):
        apply_batch(db,p,reason='过期确认',allow_overwrite=True)
    with pytest.raises(ValueError,match='重开'):
        generate(db,ids,session_ids=['S1','S2'],source_scope_confirmed=True)
    assert assignments(db)==before


def test_m1_decision_is_flagged_after_imported_session_change(db):
    seed(db)
    rid = records(db,'orders')[0]['id']
    old = assignments(db)[rid]
    old.pop('evidence')
    db.execute('UPDATE assignments SET data=? WHERE record_id=?',(dumps(old),rid))
    result = import_workbook(db,workbook_bytes({'场次':[['S1','甲','2026-09-01 09:00','2026-09-01 11:00']]}),'更正.xlsx',confirm_updates=True,reason='运营补正时间')
    assert not result['issues']
    assert assignments(db)[rid]['targets']==old['targets']
    assert assignments(db)[rid]['review_reason']
    assert db.execute("SELECT COUNT(*) FROM changes WHERE object_type='assignment_review'").fetchone()[0]>0
    assert any('需复核' in x for x in preview_close(db,'2026-09')['issues'])


def test_ad_mapping_change_invalidates_preview_and_adopted_evidence(db):
    seed(db)
    rid = records(db,'ads')[0]['id']
    save_ad_link(db,'2026-09','账号1','计划1',['S1','S2'],reason='共同受益')
    generate(db,[rid],month='2026-09')
    p = prepare_assignments(db,[rid],use_candidates=True,allow_overwrite=True)
    save_ad_link(db,'2026-09','账号1','计划1',['S1'],reason='范围变更')
    with pytest.raises(ValueError,match='依据已变化'):
        apply_batch(db,p,reason='过期预览',allow_overwrite=True)
    save_ad_link(db,'2026-09','账号1','计划1',['S1','S2'],reason='共同受益')
    p = prepare_assignments(db,[rid],use_candidates=True,allow_overwrite=True)
    apply_batch(db,p,reason='人工确认时长建议',allow_overwrite=True)
    save_ad_link(db,'2026-09','账号1','计划1',['S1'],reason='真实范围核实')
    assert fee_states(db)[rid][1]=='需复核'
    assert any('需复核' in x for x in preview_close(db,'2026-09')['issues'])


def test_cost_standard_staleness_and_explicit_actual_confirmation(db):
    seed(db)
    cost = records(db,'fulfillment')[0]
    db.execute('UPDATE records SET data=? WHERE id=?',(dumps(cost['data']|{'unit_cost':None}),cost['id']))
    sid = save_standard(db,'学习机','2026-09','2000',reason='本月采购确认')
    apply_batch(db,prepare_cost_fill(db,sid),reason='确认补齐')
    save_standard(db,'学习机','2026-09','2100',reason='标准调整')
    assert fee_states(db)[cost['id']][1]=='需复核'
    current = records(db,'fulfillment')[0]
    raw = dict(current['data'])
    raw['unit_cost'] = '2000.00'  # Formatting alone does not discard the standard source.
    for key,_,typ in SCHEMAS['fulfillment'][1]:
        if typ=='optional' and type(raw[key]) is int:
            raw[key] = str(Decimal(raw[key])/100)
    save_cost(db,'fulfillment',raw,record_id=cost['id'],reason='只补充费用说明')
    assert records(db,'fulfillment')[0]['data']['_cost_source']
    assert any('SKU成本标准' in x for x in preview_close(db,'2026-09')['issues'])
    save_cost(db,'fulfillment',raw,record_id=cost['id'],reason='该订单实际单价仍为2000',confirm_unit_cost=True)
    assert '_cost_source' not in records(db,'fulfillment')[0]['data']
    assert preview_close(db,'2026-09')['results']['S1']['final_profit']==112000


@pytest.mark.parametrize('value,expected',[('0',0),('不适用','不适用'),('12.34',1234)])
def test_actual_fee_batch_preserves_other_fields_and_zero_na(db,value,expected):
    seed(db)
    costs = records(db,'fulfillment')
    with pytest.raises(ValueError,match='已有金额'):
        prepare_cost_edit(db,[r['id'] for r in costs],'insurance',value)
    p = prepare_cost_edit(db,[r['id'] for r in costs],'insurance',value,allow_overwrite=True)
    apply_batch(db,p,reason='逐笔核实相同实际金额',allow_overwrite=True)
    for before,after in zip(costs,records(db,'fulfillment')):
        assert after['data']==before['data']|{'insurance':expected}


@pytest.mark.parametrize('value',[0,-101])
def test_reused_weights_preserve_zero_negative_and_tail_cents(db,value):
    seed(db)
    monthly = records(db,'monthly')
    assign(db,monthly[0]['id'],{'S1':1,'S2':2},basis='受益权重',reason='确认')
    target = monthly[-1]
    db.execute('UPDATE records SET data=? WHERE id=?',(dumps(target['data']|{'amount':value}),target['id']))
    p = prepare_allocation_reuse(db,monthly[0]['id'],[target['id']])
    assert sum(p['items'][0]['allocation'].values())==value
    apply_batch(db,p,reason='核对受益范围',allow_overwrite=True)
    assert assignments(db)[target['id']]['targets'] == p['items'][0]['allocation']


def test_recommendation_history_and_batch_id_in_closed_export(db):
    ids = unassign(db)
    generate(db,ids,session_ids=['S1','S2'])  # Preserve earlier exceptions too.
    generate(db,ids,session_ids=['S1','S2'],source_scope_confirmed=True)
    p = prepare_assignments(db,ids,use_candidates=True)
    apply_batch(db,p,reason='运营范围核实')
    snapshot = close_month(db,'2026-09',confirmed=True)
    assert len(snapshot['recommendations'])==4
    wb = load_workbook(BytesIO(export_management(db,'2026-09')[1]))
    assert len(wb.sheetnames)==7
    assert any(p['operation_id'] in row for row in wb['人工修改'].values)
    assert any('尚未确认' in str(v) for row in wb['核对明细'].values for v in row)
    assert all(c.data_type!='f' for s in wb for row in s for c in row)


def test_ad_link_session_scope_exports_as_readable_text(db):
    seed(db)
    save_ad_link(db,'2026-09','账号1','计划1',['S1','S2'],reason='两场共同受益')
    close_month(db,'2026-09',confirmed=True)
    wb = load_workbook(BytesIO(export_management(db,'2026-09')[1]))
    rows = list(wb['人工修改'].values)
    scope_rows = [row for row in rows if row[2] == '受益场次范围']
    assert scope_rows
    assert scope_rows[-1][4] == 'S1、S2'


def test_v1_backup_migration_preserves_close_and_export(tmp_path):
    path = tmp_path/'m1.db'
    conn = connect(path)
    seed(conn)
    snapshot = close_month(conn,'2026-09',confirmed=True)
    for key in ('recommendations','cost_standards','ad_links'):
        snapshot.pop(key)
    for change in snapshot['changes']:
        change.pop('operation_id',None)
    conn.execute('UPDATE closes SET snapshot=?',(dumps(snapshot),))
    before = export_management(conn,'2026-09')[1]
    for table in ('recommendations','cost_standards','ad_links'):
        conn.execute(f'DROP TABLE {table}')
    conn.execute('DROP INDEX changes_operation')
    conn.execute('ALTER TABLE changes DROP COLUMN operation_id')
    conn.execute('PRAGMA user_version=1')
    conn.close()
    upgraded = connect(path)
    assert upgraded.execute('PRAGMA user_version').fetchone()[0]==3
    assert get_close(upgraded,'2026-09')==snapshot
    after = export_management(upgraded,'2026-09')[1]
    old,new = load_workbook(BytesIO(before)),load_workbook(BytesIO(after))
    assert {s.title:list(s.values) for s in old}=={s.title:list(s.values) for s in new}
    backups = list(tmp_path.glob('m1.db.pre-m2-*.bak'))
    assert len(backups)==1
    with sqlite3.connect(backups[0]) as backup:
        assert backup.execute('PRAGMA user_version').fetchone()[0]==1
        assert json.loads(backup.execute('SELECT snapshot FROM closes').fetchone()[0])==snapshot
    upgraded.close()
    connect(path).close()
    assert len(list(tmp_path.glob('*.bak')))==1


def test_failed_upgrade_is_atomic_and_keeps_backup(tmp_path):
    path = tmp_path/'broken.db'
    with sqlite3.connect(path) as conn:
        conn.executescript('CREATE TABLE records(id INTEGER PRIMARY KEY); CREATE TABLE cost_standards(id INTEGER); PRAGMA user_version=1;')
    with pytest.raises(sqlite3.OperationalError):
        connect(path)
    with sqlite3.connect(path) as conn:
        assert conn.execute('PRAGMA user_version').fetchone()[0]==1
        assert conn.execute("SELECT name FROM sqlite_master WHERE name='recommendations'").fetchone() is None
    assert len(list(tmp_path.glob('*.bak')))==1


def test_payment_time_minute_filter_and_known_business_month(db):
    seed(db)
    rows,_ = order_queue(db,{'order_status':'all','time_from':'2026-09-01 10:30','time_to':'2026-09-01 10:30'})
    assert [r['data']['order_id'] for r in rows]==['O1']
    rows,_ = order_queue(db,{'order_status':'all','time_from':'11:00','time_to':'11:30'})
    assert [r['data']['order_id'] for r in rows]==['O2']
    rows,_ = order_queue(db,{'order_status':'all','time_to':'10:30'})
    assert [r['data']['order_id'] for r in rows]==['O1']
    assert order_queue(db,{'order_status':'all','month':'2026-10'})[0]==[]
    with pytest.raises(ValueError,match='截止'):
        order_queue(db,{'time_from':'2026-09-01 11:00','time_to':'2026-09-01 10:30'})
    with pytest.raises(ValueError,match='HH:MM'):
        order_queue(db,{'time_to':'25:30'})
