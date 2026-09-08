from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.workbench.db import connect, records
from src.workbench.imports import import_workbook, make_template, preview_workbook


@pytest.fixture
def db():
    conn = connect(':memory:')
    yield conn
    conn.close()


def workbook_bytes(rows=None):
    wb = load_workbook(BytesIO(make_template()))
    for sheet, entries in (rows or {}).items():
        for entry in entries:
            wb[sheet].append(entry)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


SESSION = ['S1', '达人甲', '2026-09-01 10:00', '2026-09-01 11:00']


def test_hash_rename_business_key_and_confirm_update(db):
    raw = workbook_bytes({'场次': [SESSION]})
    assert import_workbook(db, raw, 'one.xlsx')['accepted'] == 1
    assert import_workbook(db, raw, 'renamed.xlsx')['accepted'] == 0
    changed = workbook_bytes({'场次': [SESSION[:1] + ['达人乙'] + SESSION[2:]]})
    preview = preview_workbook(db, changed)
    assert len(preview['updates']) == 1
    assert import_workbook(db, changed, 'update.xlsx')['accepted'] == 0
    assert records(db)[0]['data']['talent'] == '达人甲'
    assert import_workbook(db, changed, 'update.xlsx', confirm_updates=True, reason='核对姓名')['accepted'] == 1
    assert records(db)[0]['data']['talent'] == '达人乙'
    assert db.execute('SELECT COUNT(*) FROM changes').fetchone()[0] == 1


def test_bad_row_atomic_and_duplicate_key(db):
    raw = workbook_bytes({'场次': [SESSION, SESSION]})
    assert import_workbook(db, raw, 'bad.xlsx')['accepted'] == 0
    assert records(db) == []
    assert db.execute('SELECT preview FROM batches').fetchone()


def test_missing_header_and_illegal_amount(db):
    wb = load_workbook(BytesIO(make_template()))
    wb['场次']['A1'] = '错误标题'
    stream = BytesIO(); wb.save(stream)
    assert preview_workbook(db, stream.getvalue())['issues']
    raw = workbook_bytes({'场次': [SESSION], '投流': [['ad1', 'account', 'plan',
        '2026-09-01 10:00', '2026-09-01 11:00', 'NaN']]})
    assert import_workbook(db, raw, 'bad.xlsx')['accepted'] == 0
    assert records(db) == []


def test_unknown_columns_and_formula_rejected(db):
    wb = load_workbook(BytesIO(make_template()))
    wb['场次']['E1'] = '客户电话'
    stream = BytesIO(); wb.save(stream)
    assert preview_workbook(db, stream.getvalue())['issues']
    raw = workbook_bytes({'场次': [['=1+1', *SESSION[1:]]]})
    assert preview_workbook(db, raw)['issues']


def test_missing_and_na_differ(db):
    raw = workbook_bytes({'场次': [SESSION], '达人费用': [['S1', None, '不适用', '0', '结算单']]})
    assert import_workbook(db, raw, 'fees.xlsx')['accepted'] == 2
    fee = records(db, 'talent')[0]['data']
    assert fee['commission'] is None and fee['slot_fee'] == '不适用' and fee['talent_adjustment'] == 0
