"""Explicit standard XLSX contract, preview and atomic import (not a platform adapter)."""

import hashlib
import json
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.engine.session_profit import decimal_text, money_to_cents
from src.workbench.db import audit, dumps, ensure_open, now, records, session_map, transaction

# (internal field, standard Chinese header, type). '?' means missing is retained.
SCHEMAS = {
    'sessions': ('场次', [('session_id', '场次编号', 'text'), ('talent', '达人', 'text'),
        ('start', '开播时间', 'time'), ('end', '结束时间', 'time')]),
    'orders': ('订单行', [('order_id', '订单号', 'text'), ('line_id', '子订单号或行号', 'text'),
        ('sku', 'SKU', 'text'), ('quantity', '数量', 'int'), ('paid_at', '支付时间', 'time'),
        ('status', '最终结算状态', 'text'), ('settled_at', '最终结算日期', 'date?'),
        ('sales', '确认结算销售额', 'money?'), ('platform_fee', '平台费', 'money?'),
        ('other_deductions', '其他未扣款', 'money?'), ('original_sales', '原始成交额', 'money?'),
        ('refund', '退款额', 'money?'), ('cancelled', '取消额', 'money?')]),
    'ads': ('投流', [('source_id', '原始记录编号', 'text'), ('account', '账号', 'text'),
        ('plan', '计划', 'text'), ('start', '小时起始', 'time'), ('end', '小时结束', 'time'),
        ('amount', '金额', 'money?')]),
    'fulfillment': ('商品与履约', [('order_id', '订单号', 'text'), ('line_id', '子订单号或行号', 'text'),
        ('unit_cost', '单位成本', 'unit?'), ('shipped', '出库数量', 'int'),
        ('restored', '退回恢复库存数量', 'int'), ('damaged', '损坏退回数量', 'int'),
        ('gift', '实际赠品费', 'optional'), ('insurance', '实际运费险', 'optional'),
        ('logistics', '实际物流费', 'optional'), ('loss', '退货损耗', 'optional')]),
    'talent': ('达人费用', [('session_id', '场次编号', 'text'), ('commission', '最终佣金', 'optional'),
        ('slot_fee', '坑位费', 'optional'), ('talent_adjustment', '达人调整', 'optional'),
        ('evidence', '结算依据', 'text?')]),
    'monthly': ('月度费用', [('source_id', '费用编号', 'text'), ('month', '月份', 'month'),
        ('category', '费用类别', 'text'), ('amount', '金额', 'optional'), ('basis', '分摊依据及直播承担范围', 'text')]),
}
MONTHLY_CATEGORIES = {'warehouse': '仓储', 'labor': '人工', 'management': '管理', 'tax': '损益税费', 'adjustment': '其他结算调整'}
NA = '不适用'


def business_key(kind, data):
    if kind in ('orders', 'fulfillment'):
        return dumps([data['order_id'], data['line_id']])
    if kind in ('sessions', 'talent'):
        return data['session_id']
    return data['source_id']


def parse_value(value, field_type):
    if value is None or (isinstance(value, str) and not value.strip()):
        if field_type.endswith('?') or field_type == 'optional':
            return None
        raise ValueError('必填字段为空')
    if isinstance(value, bool):
        raise ValueError('不能使用布尔值')
    if isinstance(value, (datetime, date)):
        value = value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    text = str(value).strip()
    kind = field_type.rstrip('?')
    if kind == 'optional' and text == NA:
        return NA
    if kind in ('money', 'optional'):
        return money_to_cents(text)
    if kind == 'unit':
        if decimal_text(text) < 0:
            raise ValueError('单位成本不能为负')
        return text
    if kind == 'int':
        if not text.isascii() or not text.isdigit():
            raise ValueError('数量必须是非负整数')
        value = int(text)
        if value > 1_000_000_000:
            raise ValueError('数量超范围')
        return value
    if kind == 'time':
        dt = datetime.fromisoformat(text)
        if dt.tzinfo:
            raise ValueError('请使用北京时间，不带时区')
        return dt.isoformat(sep=' ', timespec='seconds')
    if kind == 'date':
        return date.fromisoformat(text[:10]).isoformat()
    if kind == 'month':
        if len(text) != 7:
            raise ValueError('月份应为YYYY-MM')
        date.fromisoformat(text + '-01')
    if len(text) > 2000:
        raise ValueError('文本过长')
    return text


def normalize(kind, raw):
    data = {}
    for field, label, typ in SCHEMAS[kind][1]:
        try:
            data[field] = parse_value(raw.get(field), typ)
        except (ValueError, TypeError, OverflowError) as exc:
            raise ValueError(f'{label}: {exc}') from exc
    if kind in ('sessions', 'ads') and data['end'] <= data['start']:
        raise ValueError('结束时间必须晚于开始时间')
    if kind == 'orders':
        if data['quantity'] <= 0:
            raise ValueError('订单数量必须为正整数')
        if data['status'] not in ('已结算', '待结算'):
            raise ValueError('最终结算状态应为已结算或待结算')
        if data['settled_at'] and data['settled_at'] < data['paid_at'][:10]:
            raise ValueError('结算日期早于支付日期')
    if kind == 'fulfillment' and data['restored'] + data['damaged'] > data['shipped']:
        raise ValueError('退回数量大于出库数量')
    if kind == 'monthly' and data['category'] not in MONTHLY_CATEGORIES.values():
        raise ValueError('费用类别应为仓储、人工、管理、损益税费或其他结算调整')
    for field, _, typ in SCHEMAS[kind][1]:
        value = data[field]
        signed = field == 'talent_adjustment' or (kind == 'monthly' and data['category'] == '其他结算调整')
        if typ in ('money?', 'optional') and type(value) is int and value < 0 and not signed:
            raise ValueError(f'{field}不能为负；调整请使用调整科目')
    return data


def make_template():
    wb = Workbook()
    wb.remove(wb.active)
    for kind, (name, fields) in SCHEMAS.items():
        ws = wb.create_sheet(name)
        ws.append([label for _, label, _ in fields])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(fields))}1000'
        for col, (_, label, typ) in enumerate(fields, 1):
            ws.column_dimensions[get_column_letter(col)].width = max(20, min(42, len(label) * 2 + 4))
            cell = ws.cell(1, col)
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='173D50')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            # Text inputs protect identifiers and preserve Decimal unit prices.
            for row in range(2, 102):
                ws.cell(row, col).number_format = '@'
            if typ == 'optional':
                rule = DataValidation(type='custom', formula1=f'OR(ISNUMBER({get_column_letter(col)}2),ISTEXT({get_column_letter(col)}2))', allow_blank=True)
                rule.promptTitle = '费用口径'
                rule.prompt = '留空=未取得；0=实际零元；不适用=该项业务不发生。'
                rule.showInputMessage = True
                ws.add_data_validation(rule)
                rule.add(f'{get_column_letter(col)}2:{get_column_letter(col)}10000')
        ws.row_dimensions[1].height = 34
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def affected_months(conn, kind, key, data, record_id=None):
    sessions = session_map(conn)
    months = set()
    if kind == 'sessions':
        months.add(data['start'][:7])
    elif kind == 'monthly':
        months.add(data['month'])
    elif kind == 'talent':
        if data['session_id'] in sessions:
            months.add(sessions[data['session_id']]['start'][:7])
    elif kind in ('orders', 'ads'):
        months.add(data['paid_at' if kind == 'orders' else 'start'][:7])
    if kind == 'fulfillment':
        order = conn.execute("SELECT id,data FROM records WHERE kind='orders' AND business_key=?", (key,)).fetchone()
        if order:
            record_id = order['id']
            months.add(json.loads(order['data'])['paid_at'][:7])
    if record_id:
        row = conn.execute('SELECT data FROM assignments WHERE record_id=?', (record_id,)).fetchone()
        if row:
            for sid in json.loads(row['data'])['targets']:
                if sid in sessions:
                    months.add(sessions[sid]['start'][:7])
    return months


def validate_references(all_records):
    sessions = {key for (kind, key) in all_records if kind == 'sessions'}
    issues = []
    for (kind, key), data in all_records.items():
        if kind == 'talent' and data['session_id'] not in sessions:
            issues.append(f'达人费用 {key}: 场次不存在')
        if kind == 'fulfillment':
            order = all_records.get(('orders', key))
            if not order:
                issues.append(f'商品与履约 {key}: 订单行不存在')
            elif data['shipped'] > order['quantity']:
                issues.append(f'商品与履约 {key}: 出库数量大于订单数量')
    return issues


def _new_session_id(source_id, data, digest, position, existing_ids):
    """Stable for one file preview; readable without trusting a source identifier."""
    stamp = data['start'].replace('-', '')[:8] + '-' + data['start'][11:16].replace(':', '')
    base = f'LS-{stamp}-{digest[:6].upper()}-{position:02d}'
    candidate, suffix = base, 1
    while candidate in existing_ids:
        suffix += 1
        candidate = f'{base}-{suffix}'
    return candidate


def preview_workbook(conn, raw, *, import_mode='legacy', selected_session_ids=None,
                     confirm_new_session=False, duplicate_basis='', finalize=False):
    if import_mode not in ('legacy', 'new', 'supplement'):
        raise ValueError('请选择新增直播场次或补充/更正已有场次')
    selected_session_ids = sorted(set(selected_session_ids or []))
    known_sessions = session_map(conn)
    if import_mode == 'supplement' and any(s not in known_sessions for s in selected_session_ids):
        raise ValueError('补充范围包含不存在的场次')
    result = {'issues': [], 'warnings': [], 'updates': [], 'rows': [], 'members': [],
              'new_count': 0, 'update_count': 0, 'unchanged': 0, 'conflict_count': 0,
              'duplicate_file': False, 'session_mappings': {},
              'context': {'import_mode': import_mode,
                          'selected_session_ids': selected_session_ids,
                          'duplicate_basis': duplicate_basis.strip()}}
    digest = hashlib.sha256(raw).hexdigest()
    result['file_hash'] = digest
    if conn.execute("SELECT 1 FROM batches WHERE file_hash=? AND status='accepted'", (digest,)).fetchone():
        result['duplicate_file'] = True
        result['issues'].append('相同文件已入账（含改名文件）')
        return result
    existing = {(r['kind'], r['business_key']): r for r in records(conn)}
    combined = {key: r['data'] for key, r in existing.items()}
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    except Exception as exc:
        result['issues'].append(f'无法读取Excel: {type(exc).__name__}')
        return result
    seen = set()
    source_session_map = {}
    generated_session_ids = set(known_sessions)
    for kind, (name, fields) in SCHEMAS.items():
        if name not in wb.sheetnames:
            if kind != 'monthly':
                result['issues'].append(f'缺少工作表: {name}')
            continue
        ws = wb[name]
        iterator = ws.iter_rows()
        header = [c.value for c in next(iterator, [])]
        while header and header[-1] is None:
            header.pop()
        expected = [label for _, label, _ in fields]
        if header != expected:
            result['issues'].append(f'{name}: 表头必须与标准模板完全一致，不导入额外字段')
            continue
        for row_no, cells in enumerate(iterator, 2):
            if all(c.value is None for c in cells):
                continue
            conflict_key, conflict_data = '', {}
            try:
                if any(c.data_type == 'f' for c in cells):
                    raise ValueError('不接受公式，请粘贴为值')
                if any(c.value is not None for c in cells[len(fields):]):
                    raise ValueError('存在模板以外的字段')
                values = {field: cells[i].value if i < len(cells) else None for i, (field, _, _) in enumerate(fields)}
                data = normalize(kind, values)
                source_session_id = None
                if kind == 'sessions' and import_mode == 'new':
                    source_session_id = data['session_id']
                    if source_session_id in source_session_map:
                        raise ValueError('文件内场次编号重复')
                    internal_id = _new_session_id(source_session_id, data, digest,
                                                  len(source_session_map) + 1,
                                                  generated_session_ids)
                    source_session_map[source_session_id] = internal_id
                    generated_session_ids.add(internal_id)
                    data = data | {'session_id': internal_id}
                    if source_session_id in known_sessions:
                        result['warnings'].append(
                            f'来源场次编号 {source_session_id} 已存在；本次按新场次处理并映射为 {internal_id}')
                elif kind == 'talent' and data['session_id'] in source_session_map:
                    data = data | {'session_id': source_session_map[data['session_id']]}
                if (import_mode == 'supplement' and selected_session_ids
                        and kind in ('sessions', 'talent')
                        and data['session_id'] not in selected_session_ids):
                    raise ValueError('记录超出本次所选场次范围；请补选场次或拆分文件')
                key = business_key(kind, data)
                conflict_key, conflict_data = key, data
                pair = (kind, key)
                if pair in seen:
                    raise ValueError('文件内业务键重复')
                seen.add(pair)
                old = existing.get(pair)
                combined[pair] = data
                if (import_mode == 'new' and kind == 'orders' and old
                        and old['data'] != data):
                    raise ValueError('新增模式发现已有订单且内容变化；请转入补充或更正已有场次')
                member = {'kind': kind, 'key': key, 'data': data, 'row': row_no,
                          'source_session_id': source_session_id}
                if old and old['data'] == data:
                    result['unchanged'] += 1
                    result['members'].append(member | {'action': 'unchanged', 'record_id': old['id']})
                    continue
                months = affected_months(conn, kind, key, data, old['id'] if old else None)
                if old:
                    months |= affected_months(conn, kind, key, old['data'], old['id'])
                ensure_open(conn, months)
                entry = {'kind': kind, 'key': key, 'data': data, 'row': row_no, 'old_id': old['id'] if old else None}
                result['rows'].append(entry)
                if old:
                    result['update_count'] += 1
                    result['updates'].append({'kind': kind, 'key': key, 'old': old['data'], 'new': data})
                    result['members'].append(member | {'action': 'update', 'record_id': old['id']})
                else:
                    result['new_count'] += 1
                    result['members'].append(member | {'action': 'new', 'record_id': None})
            except (ValueError, TypeError) as exc:
                result['issues'].append(f'{name} 第{row_no}行: {exc}')
                result['conflict_count'] += 1
                result['members'].append({'kind': kind, 'key': conflict_key,
                                          'data': conflict_data, 'row': row_no,
                                          'action': 'conflict', 'record_id': None,
                                          'error': str(exc)})
    wb.close()
    result['session_mappings'] = source_session_map
    if import_mode == 'new' and not source_session_map:
        result['issues'].append('新增直播场次模式必须在“场次”工作表提供至少一场新直播')
    if import_mode == 'new':
        for source_id, internal_id in source_session_map.items():
            session = combined.get(('sessions', internal_id))
            if not session:
                continue
            similar = [sid for sid, old in known_sessions.items()
                       if old['talent'] == session['talent'] and old['start'] == session['start']
                       and old['end'] == session['end']]
            if similar:
                result['warnings'].append(
                    f'{source_id} 与已有场次 {"、".join(similar)} 的达人和时间相同，请核对是否重复')
        if finalize and result['warnings'] and not confirm_new_session:
            result['issues'].append('发现来源编号冲突或疑似重复场次；确认确为新场次并填写依据后才能入账')
        if finalize and result['warnings'] and confirm_new_session and not duplicate_basis.strip():
            result['issues'].append('确认新场次时必须填写编号冲突或并行直播的核对依据')
    result['issues'].extend(validate_references(combined))
    return result


def import_workbook(conn, raw, filename, *, confirm_updates=False, reason='',
                    import_mode='legacy', selected_session_ids=None,
                    confirm_new_session=False, duplicate_basis=''):
    with transaction(conn):
        preview = preview_workbook(
            conn, raw, import_mode=import_mode,
            selected_session_ids=selected_session_ids,
            confirm_new_session=confirm_new_session,
            duplicate_basis=duplicate_basis, finalize=True)
        if preview['updates'] and (confirm_updates is not True or not reason.strip()):
            preview['issues'].append('同键内容变化，需核对差异并填写原因后确认更新')
        status = 'rejected' if preview['issues'] else 'accepted'
        batch = conn.execute(
            'INSERT INTO batches(filename,file_hash,created_at,status,preview,import_mode,context) '
            'VALUES(?,?,?,?,?,?,?)',
            (Path(filename).name, preview['file_hash'], now(), status, dumps(preview),
             import_mode, dumps(preview['context']))).lastrowid
        count = 0
        resolved_ids = {}
        if not preview['issues']:
            for entry in preview['rows']:
                if entry['old_id']:
                    old = conn.execute('SELECT * FROM records WHERE id=?', (entry['old_id'],)).fetchone()
                    audit(conn, entry['kind'], entry['key'], dict(old),
                          {'data': entry['data'], 'batch_id': batch, 'source_row': entry['row']}, reason)
                    conn.execute('UPDATE records SET data=?,batch_id=?,source_row=? WHERE id=?',
                        (dumps(entry['data']), batch, entry['row'], entry['old_id']))
                    resolved_ids[(entry['kind'], entry['key'])] = entry['old_id']
                    flag_changed_assignments(conn, entry['kind'], entry['key'], json.loads(old['data']), entry['data'], entry['old_id'])
                    # Keep the decision and its original source amount. Closing detects stale amounts.
                else:
                    identifier = conn.execute(
                        'INSERT INTO records(kind,business_key,data,batch_id,source_row) VALUES(?,?,?,?,?)',
                        (entry['kind'], entry['key'], dumps(entry['data']), batch, entry['row'])).lastrowid
                    resolved_ids[(entry['kind'], entry['key'])] = identifier
                count += 1
            for source_id, internal_id in preview['session_mappings'].items():
                conn.execute(
                    'INSERT INTO session_source_mappings(batch_id,source_session_id,internal_session_id,created_at) '
                    'VALUES(?,?,?,?)', (batch, source_id, internal_id, now()))
        for position, member in enumerate(preview['members'], 1):
            rid = member.get('record_id')
            if status == 'accepted':
                rid = rid or resolved_ids.get((member['kind'], member['key']))
            conn.execute(
                'INSERT INTO batch_members(batch_id,member_order,kind,business_key,action,record_id,source_row,snapshot) '
                'VALUES(?,?,?,?,?,?,?,?)',
                (batch, position, member['kind'], member['key'], member['action'], rid,
                 member.get('row'), dumps(member)))
        return {'batch_id': batch, 'accepted': count, **preview}


def batch_member_ids(conn, batch_id, kind=None, actions=None):
    sql = 'SELECT record_id FROM batch_members WHERE batch_id=? AND record_id IS NOT NULL'
    params = [int(batch_id)]
    if kind:
        sql += ' AND kind=?'
        params.append(kind)
    if actions:
        actions = list(actions)
        sql += ' AND action IN (' + ','.join('?' for _ in actions) + ')'
        params.extend(actions)
    return [r['record_id'] for r in conn.execute(sql + ' ORDER BY member_order', params)]


def recent_batches(conn, limit=20):
    rows = conn.execute(
        "SELECT b.*, SUM(CASE WHEN m.action='new' THEN 1 ELSE 0 END) AS new_count, "
        "SUM(CASE WHEN m.action='update' THEN 1 ELSE 0 END) AS update_count, "
        "SUM(CASE WHEN m.action='unchanged' THEN 1 ELSE 0 END) AS unchanged_count, "
        "SUM(CASE WHEN m.action='conflict' THEN 1 ELSE 0 END) AS conflict_count "
        "FROM batches b LEFT JOIN batch_members m ON m.batch_id=b.id "
        "GROUP BY b.id ORDER BY b.id DESC LIMIT ?", (int(limit),))
    return [dict(r) | {'context': json.loads(r['context'])} for r in rows]


def flag_changed_assignments(conn, kind, key, old, new, record_id):
    """M1 decisions have no evidence copy: flag future changes without inventing old evidence."""
    fields = {'orders':('paid_at',), 'ads':('account','plan','start','end','amount'),
              'monthly':('month','category','amount'), 'sessions':('start','end','talent')}.get(kind,())
    if not any(old.get(k) != new.get(k) for k in fields):
        return
    for row in conn.execute('SELECT * FROM assignments').fetchall():
        decision = json.loads(row['data'])
        if row['record_id'] == record_id or (kind == 'sessions' and key in decision['targets']):
            updated = decision | {'review_reason':'来源或场次关键依据已变更，请人工复核'}
            audit(conn,'assignment_review',row['record_id'],decision,updated,'来源或场次更正引发归属/分配复核')
            conn.execute('UPDATE assignments SET data=? WHERE record_id=?',(dumps(updated),row['record_id']))
