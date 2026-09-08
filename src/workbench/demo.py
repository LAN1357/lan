"""Generate a standard, explicitly synthetic two-session workbook for manual trial."""

import argparse
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from src.workbench.imports import make_template

DEMO_ROWS = {
    '场次': [['S1', '合成达人甲', '2026-09-01 10:00', '2026-09-01 11:00'],
             ['S2', '合成自播乙', '2026-09-01 11:00', '2026-09-01 12:00']],
    '订单行': [['O1', '1', '合成学习机', 2, '2026-09-01 10:30', '已结算', '2026-10-01', '9000', '300', '0', '10000', '1000', '0'],
              ['O2', '1', '合成学习机', 3, '2026-09-01 11:30', '已结算', '2026-09-10', '3000', '100', '0', '9000', '6000', '0']],
    '商品与履约': [['O1', '1', '2000', 2, 0, 0, '100', '30', '100', '50'],
                   ['O2', '1', '500', 3, 1, 1, '不适用', '0', '20', '500']],
    '达人费用': [['S1', '900', '500', '0', '合成最终结算单']],
    '投流': [['A1', '合成账号', '合成计划', '2026-09-01 10:30', '2026-09-01 11:30', '1200']],
    '月度费用': [['W', '2026-09', '仓储', '100', '合成账例：直播承担部分'],
                 ['L', '2026-09', '人工', '300', '合成账例：直播承担部分'],
                 ['M', '2026-09', '管理', '200', '合成账例：直播承担部分'],
                 ['T', '2026-09', '损益税费', '150', '合成账例：BP损益确认'],
                 ['J', '2026-09', '其他结算调整', '50', '合成结算调整']],
}


def demo_workbook():
    wb = load_workbook(BytesIO(make_template()))
    for name, rows in DEMO_ROWS.items():
        ws = wb[name]
        # The template reserves text-formatted input rows. Write from row 2, not append.
        for row_no, row in enumerate(rows, 2):
            for col, value in enumerate(row, 1):
                ws.cell(row_no, col, value)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def main():
    parser = argparse.ArgumentParser(description='只生成合成标准模板，不修改数据库')
    parser.add_argument('--output', type=Path, default=Path('docs/templates'))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    for name, raw in [('经营复盘标准模板.xlsx', make_template()), ('经营复盘合成两场账例.xlsx', demo_workbook())]:
        path = args.output / name
        path.write_bytes(raw)
        print(path.resolve())


if __name__ == '__main__':
    main()
