"""Formal XLSX exports read only immutable close copies, never live amounts."""

from decimal import Decimal
import json
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.workbench.closing import get_close
from src.workbench.imports import SCHEMAS

LABELS = dict(sales='确认结算销售额', platform_fee='平台费', other_deductions='其他未扣款',
    commission='最终佣金', slot_fee='坑位费', talent_adjustment='达人调整', ad_spend='投流费',
    product_cost='商品销售成本', gift='赠品费', insurance='运费险', logistics='物流费', loss='退货损耗',
    warehouse='仓储分摊', labor='人工分摊', management='管理分摊', tax='损益税费分摊', adjustment='其他结算调整',
    net_revenue='净经营收入', direct_cost='直接成本合计', indirect_cost='间接费用及税费合计',
    operating_profit='经营复盘利润', final_profit='最终结算利润', operating_margin='经营复盘利润率',
    final_margin='最终结算利润率', ad_roi='结算销售额/投流费')
MONEY_FORMAT = '#,##0.00;[Red](#,##0.00);0.00'


def yuan(cents):
    return Decimal(cents) / 100


def change_values(raw):
    value = json.loads(raw) if raw else None
    if isinstance(value, dict) and 'data' in value:
        value = value['data']
        if isinstance(value, str):
            value = json.loads(value)
    return value if isinstance(value, dict) else {}


def change_cell(value, key, kind):
    if value is None:
        return '未记录', ''
    if type(value) is bool:
        return ('适用' if value else '不适用'), ''
    if isinstance(value, (list, tuple)):
        return ('、'.join(str(item) for item in value) or '无'), ''
    if key=='evidence' and isinstance(value,dict):
        labels = {k:label for _,fields in SCHEMAS.values() for k,label,_ in fields}
        source = '；'.join(f'{labels.get(k,k)}：{str(yuan(v))+" 元" if k=="amount" and type(v) is int else v}' for k,v in value.get('source',{}).items())
        sessions = '；'.join(f'{s} / {d["talent"]} / {d["start"]}—{d["end"]}' for s,d in value.get('sessions',{}).items() if d)
        return source + '；目标场次：' + sessions, ''
    if key=='_cost_source' and isinstance(value,dict):
        return f'{value["sku"]} / {value["month"]} / {value["data"]["unit_cost"]} 元；{value["data"]["basis"]}', ''
    if key in ('source_amount',) or any(f == key and typ in ('money?', 'optional') for f, _, typ in SCHEMAS.get(kind, ('', []))[1]):
        return (yuan(value), '元') if type(value) is int else (value, '元')
    if isinstance(value, dict):
        return '；'.join(f'{k}: {v}' for k, v in value.items()), '分（费用）/归属标记（订单）' if key == 'targets' else ('分' if key == 'tail_cents' else '')
    return value, ''


def append(ws, values):
    # Closed snapshots contain a few structured audit fields. Keep exports
    # readable and never pass containers through to openpyxl cells.
    values = [json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list, tuple)) else value for value in values]
    ws.append(values)
    # Even untrusted strings starting with '=' are literal text in the saved XLSX.
    for cell, value in zip(ws[ws.max_row], values):
        if isinstance(value, str):
            cell.data_type = 's'


def style(ws, money_columns=(), widths=None, *, header_row=1):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'B{header_row + 1}'
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}'
    ws.print_title_rows = f'1:{header_row}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws.max_column > 8 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_area = ws.dimensions
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    for row in ws:
        for cell in row:
            cell.font = Font(name='Arial', size=11, color='243746')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row[0].row > header_row and row[0].row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F2F5F6')
            if cell.column in money_columns and cell.row > header_row and cell.data_type == 'n':
                cell.number_format = MONEY_FORMAT
        ws.row_dimensions[row[0].row].height = 32
    for cell in ws[header_row]:
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='173D50')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[header_row].height = 36
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 21)


def export_management(conn, month, *, version=None):
    snapshot = get_close(conn, month, version)
    if snapshot is None:
        raise ValueError('仅能导出有效已关账版本；重开期间请明确选择历史版本')
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet('管理摘要')
    append(summary, ['月度经营复盘', month, f'V{snapshot["version"]}'])
    append(summary, ['金额单位：人民币元；同一开播业务月份', None, None])
    append(summary, ['指标', '本月合计', '口径说明'])
    summary_keys = ['sales', 'net_revenue', 'direct_cost', 'operating_profit', 'indirect_cost', 'adjustment', 'final_profit', 'operating_margin', 'final_margin']
    for key in summary_keys:
        value = snapshot['total'][key]
        percentage = key.endswith('margin')
        append(summary, [LABELS[key], ('不适用' if value is None else Decimal(value)) if percentage else yuan(value),
                         '分母为净经营收入；非正时不适用' if percentage else ''])
        if percentage and value is not None:
            summary.cell(summary.max_row, 2).number_format = '0.0%'
    style(summary, widths={1: 34, 2: 24, 3: 47}, header_row=3)
    for row in range(4, 11):
        summary.cell(row, 2).number_format = MONEY_FORMAT
    for row in (7, 10):
        for cell in summary[row]:
            cell.font = Font(name='Arial', size=11, bold=True)

    # A vertical P&L stays readable and printable even with all 17 fixed report categories.
    pnl = wb.create_sheet('场次损益')
    session_ids = sorted(snapshot['results'])
    append(pnl, ['科目（人民币元）'] + [f'{sid} / {snapshot["sessions"][sid]["talent"]}' for sid in session_ids] + ['合计'])
    for key in LABELS:
        ratio = key.endswith('margin') or key == 'ad_roi'
        vals = [snapshot['results'][sid][key] for sid in session_ids] + [snapshot['total'][key]]
        append(pnl, [LABELS[key]] + [('不适用' if v is None else Decimal(v)) if ratio else yuan(v) for v in vals])
        for cell in pnl[pnl.max_row][1:]:
            cell.number_format = ('0.0%' if key.endswith('margin') else '0.000') if ratio else MONEY_FORMAT
    style(pnl, widths={1: 32})

    detail = wb.create_sheet('核对明细')
    append(detail, ['来源类型', '业务键', '归属场次', '字段', '采用值', '单位', '来源文件', '批次', '来源行', '跨期说明'])
    batches = {b['id']: b for b in snapshot['batches']}
    source_rows = {str(r['id']): r for r in snapshot['records']}
    orders_by_key = {r['business_key']:r for r in snapshot['records'] if r['kind']=='orders'}
    for row in snapshot['records']:
        kind, data = row['kind'], row['data']
        assigned = snapshot['assignments'].get(str(row['id']), {}).get('targets', {})
        if kind == 'fulfillment':
            order = orders_by_key.get(row['business_key'])
            assigned = snapshot['assignments'].get(str(order['id']), {}).get('targets', {}) if order else {}
        session_text = ', '.join(assigned) or data.get('session_id', '')
        source_name = batches.get(row['batch_id'], {}).get('filename', '人工维护')
        cross = ''
        if kind == 'orders':
            if data['paid_at'][:7] != month or data['settled_at'][:7] != month:
                cross = f'开播业务月{month}；支付月{data["paid_at"][:7]}；结算月{data["settled_at"][:7]}'
        for field, label, typ in SCHEMAS[kind][1]:
            value = data[field]
            unit = '元' if typ in ('money?', 'optional', 'unit?') else ('件' if typ == 'int' else '')
            if typ in ('money?', 'optional') and type(value) is int:
                value = yuan(value)
            elif value is None:
                value = '未取得（核对字段或已声明不适用项目）'
            append(detail, [SCHEMAS[kind][0], row['business_key'], session_text, label, value, unit,
                            source_name, row['batch_id'], row['source_row'], cross])
        origin = data.get('_cost_source')
        if origin:
            append(detail, ['已确认成本来源',row['business_key'],session_text,'采用标准及适用依据',
                f'{origin["sku"]} / {origin["month"]} / {origin["data"]["unit_cost"]} 元；{origin["data"]["basis"]}',
                '',source_name,row['batch_id'],row['source_row'],''])
    for rec in snapshot.get('recommendations',[]):
        source = source_rows[str(rec['record_id'])]
        d = rec['data']
        adopted = snapshot['assignments'].get(str(rec['record_id']),{}).get('recommendation_id') == rec['id']
        evidence_labels = {k:label for k,label,_ in SCHEMAS[source['kind']][1]}
        source_evidence = '；'.join(f'{evidence_labels.get(k,k)}：{str(yuan(v))+" 元" if k=="amount" and type(v) is int else v}' for k,v in d['evidence']['source'].items())
        values = [('记录及规则版本',f'{rec["id"]} / {d["rule_version"]} / {rec["created_at"]}'),
                  ('命中规则','；'.join(d['hits'])),('候选场次',', '.join(d['candidates']) or '无'),
                  ('异常及确认情况',('；'.join(d['issues']) or '无规则异常') + ('；关账决定采用此候选' if adopted else '；非本版采用候选')),
                  ('当时来源依据',source_evidence),
                  ('当时场次范围','；'.join(f'{s} / {v["talent"]} / {v["start"]}—{v["end"]}' for s,v in d['evidence']['scope_sessions'].items() if v)),
                  ('范围确认', 'BP已确认订单业务范围' if d['parameters'].get('source_scope_confirmed') else '未确认订单范围或使用投流精确关联')]
        link = d['evidence'].get('link')
        if link:
            values.append(('当时投流关联依据',f'{link["account"]} / {link["plan"]} / {link["month"]}；{link["data"]["basis"]}'))
        for label,value in values:
            append(detail,['规则候选',source['business_key'],'',label,value,'','候选生成时保存','','',''])
    style(detail, money_columns=(5,), widths={2: 26, 4: 28, 5: 42, 7: 32, 8: 10, 9: 10, 10: 58})
    for row in range(2, detail.max_row + 1):
        detail.row_dimensions[row].height = 46

    allocation = wb.create_sheet('分摊明细')
    append(allocation, ['来源类型', '业务键', '源额（元）', '目标场次', '分配额（元）', '权重', '尾差（分）', '依据', '尾差规则'])
    for rid, decision in snapshot['assignments'].items():
        if decision['mode'] == 'order':
            continue
        source = source_rows[rid]
        for sid, value in decision['targets'].items():
            append(allocation, [SCHEMAS[source['kind']][0], source['business_key'], yuan(decision['source_amount']), sid,
                yuan(value), (decision.get('weights') or {}).get(sid, '明确金额'), decision.get('tail_cents', {}).get(sid, 0),
                decision['basis'], decision['tail_rule']])
    style(allocation, money_columns=(3, 5), widths={8: 48, 9: 64})
    for row in range(2, allocation.max_row + 1):
        allocation.row_dimensions[row].height = 58

    changes = wb.create_sheet('人工修改')
    is_m2 = 'recommendations' in snapshot
    append(changes, ['对象', '业务键或记录号', '字段', '修改前', '修改后', '单位', '原因', '时间（UTC）'] + (['批量操作编号'] if is_m2 else []))
    change_labels = {f: label for _, fields in SCHEMAS.values() for f, label, _ in fields}
    change_labels.update(targets='目标场次及分配结果', weights='权重', mode='分配方式', basis='依据',
        source_amount='来源金额', tail_cents='尾差', tail_rule='尾差规则', talent='达人合作', ads='投流',
        slot='坑位费', gift='赠品', status='状态', version='版本',
        recommendation_id='采用候选记录号', review_reason='需复核原因', _cost_source='采用成本标准', session_ids='受益场次范围')
    object_labels = {'assignment':'人工归属及分配', 'applicability':'业务范围', 'reopen':'整月重开',
                     'assignment_review':'归属依据变化', 'cost_standard':'SKU月度成本标准', 'ad_link':'投流关联范围'}
    for row in snapshot['changes']:
        kind = row['object_type']
        before, after = change_values(row['old_value']), change_values(row['new_value'])
        for key in sorted(before.keys() | after.keys()):
            if before.get(key) == after.get(key):
                continue
            old_value, unit = change_cell(before.get(key), key, kind)
            new_value, new_unit = change_cell(after.get(key), key, kind)
            append(changes, [object_labels.get(kind, SCHEMAS.get(kind, (kind,))[0]), row['object_key'],
                '确认时依据' if key=='evidence' and kind in ('assignment','assignment_review') else change_labels.get(key,key), old_value, new_value, new_unit or unit, row['reason'], row['created_at']] + ([row.get('operation_id') or '单笔操作'] if is_m2 else []))
    style(changes, money_columns=(4,5), widths={2: 26, 3: 26, 4: 58, 5: 58, 6: 32, 7: 42, 8: 32})
    for row in range(2, changes.max_row + 1):
        changes.row_dimensions[row].height = 64
        if is_m2 and changes.cell(row,3).value=='采用候选记录号':
            changes.cell(row,4).number_format = '0'
            changes.cell(row,5).number_format = '0'

    product = wb.create_sheet('SKU与公共项目')
    append(product, ['场次', 'SKU或公共项目', '直接贡献或公共项目损益（元）', '口径'])
    for row in snapshot['sku']:
        append(product, [row['session_id'], row['sku'], yuan(row['contribution']), '仅收入减商品及履约成本；非SKU全成本利润'])
    for sid, value in snapshot['public'].items():
        append(product, [sid, '公共项目', yuan(value), '达人、投流、间接费用及税费、其他结算调整'])
        append(product, [sid, '场次合计', yuan(snapshot['results'][sid]['final_profit']), 'SKU直接贡献合计 + 公共项目 = 最终结算利润'])
    style(product, money_columns=(3,), widths={2: 28, 3: 36, 4: 65})

    version_sheet = wb.create_sheet('版本与口径')
    append(version_sheet, ['项目', '说明'])
    for row in [
        ['业务月份', month], ['关账版本', f'V{snapshot["version"]}'], ['关账及数据截止时间（UTC）', snapshot['closed_at']],
        ['计算规则', snapshot['rule_version']], ['费用清单确认', snapshot['confirmation']],
        ['输入范围', '仅支持标准模板；未声称适配抖音或千川原生导出格式'],
        ['金额精度', 'Decimal处理文本和单价，人民币整数分存储汇总，ROUND_HALF_UP；报表为关账时数值副本'],
        ['净经营收入', '确认结算销售额－平台费－其他未扣款。销售额已扣退款和取消，不再重复扣除；不能填平台净到账金额'],
        ['经营复盘利润', '净经营收入－达人最终费用－投流－商品及履约直接成本'],
        ['最终结算利润', '经营复盘利润－间接费用及损益税费分摊＋其他结算调整'],
        ['税费', '由BP确认损益口径，不计算应纳税额，不将所有税款支付当作费用'],
        ['退货', '计成本数量=出库－恢复库存退回－损坏退回；损坏计入独立损耗。部分退款未退货不减少数量'],
        ['时间', '按开播月份复盘；订单支付时间用于归属；结算状态及日期用于关账，跨期见核对明细'],
        ['分摊', '源额只含直播业务应承担部分，BP指定目标和权重，固定尾差规则，正负源额均守恒'],
        ['SKU', 'SKU直接贡献与公共项目分列，合计与场次最终利润相等，不宣称SKU全成本利润'],
        ['比率', '利润率分母为净经营收入；投流比率为结算销售额/投流费，分母非正显示不适用；不代表因果增量'],
        ['历史', '本文件仅使用指定关账副本，重开或实时数据更正不会改变本版金额'],
    ]:
        append(version_sheet, row)
    style(version_sheet, widths={1: 34, 2: 112})
    for row in range(2, version_sheet.max_row + 1):
        version_sheet.row_dimensions[row].height = 44
    stream = BytesIO()
    wb.save(stream)
    return f'经营复盘_{month}_V{snapshot["version"]}.xlsx', stream.getvalue()


def save_management(conn, month, directory, *, version=None):
    filename, raw = export_management(conn, month, version=version)
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / filename
    temporary = path.with_suffix('.xlsx.tmp')
    temporary.write_bytes(raw)
    temporary.replace(path)
    return path
